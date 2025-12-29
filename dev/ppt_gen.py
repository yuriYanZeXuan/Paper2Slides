"""
PPT Generation Demo with python-pptx
Features:
- Multiple fonts (Chalkboard SE, Arial, Helvetica Neue, etc.)
- Various colors and text weights
- Shape boxes (rounded corners, shadows)
- Image content
- Export to PPTX and PDF
"""

import subprocess
import platform
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import nsmap
from lxml import etree


class PPTGenerator:
    """PowerPoint generator with rich styling support"""
    
    def __init__(self, width: float = 16, height: float = 9):
        """Initialize presentation with custom dimensions (in inches)"""
        self.prs = Presentation()
        self.prs.slide_width = Inches(width)
        self.prs.slide_height = Inches(height)
        self.width = width
        self.height = height
        
        # color palette
        self.colors = {
            "primary": "#2563EB",      # blue
            "secondary": "#7C3AED",    # purple
            "accent": "#F59E0B",       # amber
            "success": "#10B981",      # green
            "danger": "#EF4444",       # red
            "dark": "#1F2937",         # gray-800
            "light": "#F3F4F6",        # gray-100
            "white": "#FFFFFF",
            "black": "#000000",
        }
        
        # font families
        self.fonts = {
            "display": "Chalkboard SE",
            "heading": "Helvetica Neue",
            "body": "Arial",
            "mono": "Menlo",
            "handwriting": "Bradley Hand",
            "elegant": "Georgia",
        }
    
    def _parse_color(self, color_str: str) -> RGBColor:
        """Parse hex color string to RGBColor"""
        hex_color = color_str.lstrip('#')
        r, g, b = (int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        return RGBColor(r, g, b)
    
    def add_slide(self):
        """Add a blank slide"""
        slide_layout = self.prs.slide_layouts[6]  # blank layout
        return self.prs.slides.add_slide(slide_layout)
    
    def add_textbox(self, slide, x: float, y: float, w: float, h: float,
                   text: str, font_name: str = "Arial", font_size: int = 24,
                   font_color: str = "#000000", bold: bool = False, 
                   italic: bool = False, align: str = "left",
                   vertical_align: str = "top"):
        """Add a textbox with rich text formatting"""
        textbox = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        tf = textbox.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        
        # vertical alignment
        anchor_map = {
            "top": MSO_ANCHOR.TOP,
            "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }
        tf.vertical_anchor = anchor_map.get(vertical_align, MSO_ANCHOR.TOP)
        
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = font_name
        p.font.size = Pt(font_size)
        p.font.color.rgb = self._parse_color(font_color)
        p.font.bold = bold
        p.font.italic = italic
        
        # horizontal alignment
        align_map = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }
        p.alignment = align_map.get(align, PP_ALIGN.LEFT)
        
        return textbox
    
    def add_multiline_text(self, slide, x: float, y: float, w: float, h: float,
                          lines: list, default_font: str = "Arial"):
        """Add multiline text with different formatting per line
        
        lines: list of dicts with keys:
            - text: str
            - font_name: str (optional)
            - font_size: int (optional)
            - font_color: str (optional)
            - bold: bool (optional)
            - italic: bool (optional)
        """
        textbox = slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        tf = textbox.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        
        for i, line_config in enumerate(lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            
            p.text = line_config.get("text", "")
            p.font.name = line_config.get("font_name", default_font)
            p.font.size = Pt(line_config.get("font_size", 24))
            p.font.color.rgb = self._parse_color(line_config.get("font_color", "#000000"))
            p.font.bold = line_config.get("bold", False)
            p.font.italic = line_config.get("italic", False)
            p.alignment = PP_ALIGN.LEFT
        
        return textbox
    
    def add_rounded_rectangle(self, slide, x: float, y: float, w: float, h: float,
                              fill_color: str = "#FFFFFF", 
                              border_color: Optional[str] = None,
                              border_width: float = 1.0,
                              corner_radius: float = 0.2,
                              shadow: bool = False):
        """Add a rounded rectangle shape with optional shadow"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        
        # fill color
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._parse_color(fill_color)
        
        # border
        if border_color:
            shape.line.color.rgb = self._parse_color(border_color)
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        
        # adjust corner radius via XML manipulation
        # python-pptx doesn't have direct API for this
        sp = shape._element
        prstGeom = sp.find('.//a:prstGeom', nsmap)
        if prstGeom is not None:
            avLst = prstGeom.find('a:avLst', nsmap)
            if avLst is None:
                avLst = etree.SubElement(prstGeom, '{%s}avLst' % nsmap['a'])
            # clear existing adjustments
            for child in list(avLst):
                avLst.remove(child)
            # add corner radius adjustment (value is in 1/100000 of shape size)
            radius_val = int(corner_radius * 100000 / min(w, h))
            gd = etree.SubElement(avLst, '{%s}gd' % nsmap['a'])
            gd.set('name', 'adj')
            gd.set('fmla', f'val {radius_val}')
        
        # add shadow effect
        if shadow:
            self._add_shadow_to_shape(shape)
        
        return shape
    
    def add_rectangle(self, slide, x: float, y: float, w: float, h: float,
                     fill_color: str = "#FFFFFF",
                     border_color: Optional[str] = None,
                     border_width: float = 1.0,
                     shadow: bool = False):
        """Add a regular rectangle shape"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._parse_color(fill_color)
        
        if border_color:
            shape.line.color.rgb = self._parse_color(border_color)
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        
        if shadow:
            self._add_shadow_to_shape(shape)
        
        return shape
    
    def add_oval(self, slide, x: float, y: float, w: float, h: float,
                fill_color: str = "#FFFFFF",
                border_color: Optional[str] = None,
                border_width: float = 1.0,
                shadow: bool = False):
        """Add an oval/circle shape"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        
        shape.fill.solid()
        shape.fill.fore_color.rgb = self._parse_color(fill_color)
        
        if border_color:
            shape.line.color.rgb = self._parse_color(border_color)
            shape.line.width = Pt(border_width)
        else:
            shape.line.fill.background()
        
        if shadow:
            self._add_shadow_to_shape(shape)
        
        return shape
    
    def _add_shadow_to_shape(self, shape, 
                            blur_radius: int = 50800,  # EMUs
                            distance: int = 38100,
                            direction: int = 2700000,  # 45 degrees in 1/60000 degrees
                            color: str = "#000000",
                            alpha: int = 40):
        """Add outer shadow effect to shape via XML"""
        sp = shape._element
        spPr = sp.find('.//p:spPr', nsmap)
        if spPr is None:
            return
        
        # create effectLst if not exists
        effectLst = spPr.find('a:effectLst', nsmap)
        if effectLst is None:
            effectLst = etree.SubElement(spPr, '{%s}effectLst' % nsmap['a'])
        
        # add outer shadow
        outerShdw = etree.SubElement(effectLst, '{%s}outerShdw' % nsmap['a'])
        outerShdw.set('blurRad', str(blur_radius))
        outerShdw.set('dist', str(distance))
        outerShdw.set('dir', str(direction))
        outerShdw.set('algn', 'tl')
        outerShdw.set('rotWithShape', '0')
        
        # shadow color with alpha
        srgbClr = etree.SubElement(outerShdw, '{%s}srgbClr' % nsmap['a'])
        srgbClr.set('val', color.lstrip('#'))
        alphaElem = etree.SubElement(srgbClr, '{%s}alpha' % nsmap['a'])
        alphaElem.set('val', f'{alpha * 1000}')  # percentage * 1000
    
    def add_image(self, slide, image_path: str, x: float, y: float,
                 width: Optional[float] = None, height: Optional[float] = None):
        """Add an image to the slide"""
        if not Path(image_path).exists():
            print(f"Warning: Image not found: {image_path}")
            return None
        
        left, top = Inches(x), Inches(y)
        
        if width and height:
            pic = slide.shapes.add_picture(
                image_path, left, top, 
                width=Inches(width), height=Inches(height)
            )
        elif width:
            pic = slide.shapes.add_picture(
                image_path, left, top, width=Inches(width)
            )
        elif height:
            pic = slide.shapes.add_picture(
                image_path, left, top, height=Inches(height)
            )
        else:
            pic = slide.shapes.add_picture(image_path, left, top)
        
        return pic
    
    def add_gradient_background(self, slide, color1: str, color2: str, angle: int = 90):
        """Add gradient background to slide (simplified - uses solid color)
        Note: python-pptx has limited gradient support, using fill shape instead
        """
        # add a full-slide rectangle as background
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), 
            Inches(self.width), Inches(self.height)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = self._parse_color(color1)
        bg_shape.line.fill.background()
        
        # move to back
        sp = bg_shape._element
        spTree = sp.getparent()
        spTree.remove(sp)
        spTree.insert(0, sp)
        
        return bg_shape
    
    def save(self, output_path: str):
        """Save the presentation to a file"""
        self.prs.save(output_path)
        print(f"Saved PPTX: {output_path}")
        return output_path
    
    def convert_to_pdf(self, pptx_path: str) -> Optional[str]:
        """Convert PPTX to PDF using LibreOffice"""
        pptx_path = Path(pptx_path)
        output_dir = pptx_path.parent
        
        system = platform.system().lower()
        
        if system == "darwin":  # macOS
            libreoffice_paths = [
                "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                "/usr/local/bin/libreoffice",
                "libreoffice",
                "soffice"
            ]
        elif system == "linux":
            libreoffice_paths = [
                "/usr/bin/libreoffice",
                "/usr/local/bin/libreoffice",
                "/snap/bin/libreoffice",
                "libreoffice",
                "soffice"
            ]
        elif system == "windows":
            libreoffice_paths = [
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                "soffice.exe",
                "libreoffice.exe"
            ]
        else:
            libreoffice_paths = ["libreoffice", "soffice"]
        
        for lo_path in libreoffice_paths:
            try:
                cmd = [
                    lo_path, "--headless", "--convert-to", "pdf",
                    "--outdir", str(output_dir), str(pptx_path)
                ]
                
                result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
                
                if result.returncode == 0:
                    pdf_name = pptx_path.stem + ".pdf"
                    pdf_path = output_dir / pdf_name
                    if pdf_path.exists():
                        print(f"Converted to PDF: {pdf_path}")
                        return str(pdf_path)
                        
            except (subprocess.SubprocessError, FileNotFoundError) as e:
                continue
        
        print("Warning: LibreOffice not found. Install LibreOffice for PDF conversion.")
        return None


def create_demo_presentation():
    """Create a demo presentation showcasing various features"""
    
    # initialize generator
    gen = PPTGenerator(width=16, height=9)
    
    # ==================== SLIDE 1: Title Slide ====================
    slide1 = gen.add_slide()
    
    # background color
    gen.add_gradient_background(slide1, "#1E3A5F", "#2C5282")
    
    # decorative shapes
    gen.add_oval(slide1, 14, -1, 3, 3, fill_color="#F59E0B", shadow=True)
    gen.add_rounded_rectangle(slide1, -0.5, 7, 4, 2.5, 
                             fill_color="#7C3AED", corner_radius=0.3, shadow=True)
    
    # title text with Chalkboard SE font
    gen.add_textbox(slide1, 1, 2.5, 14, 2,
                   text="Paper2Slides Demo",
                   font_name="Chalkboard SE", font_size=72,
                   font_color="#FFFFFF", bold=True, align="center")
    
    # subtitle with different font
    gen.add_textbox(slide1, 1, 4.8, 14, 1,
                   text="Showcase of python-pptx Features",
                   font_name="Helvetica Neue", font_size=32,
                   font_color="#F3F4F6", italic=True, align="center")
    
    # author info
    gen.add_textbox(slide1, 1, 6.5, 14, 0.8,
                   text="Created with Python • Powered by python-pptx",
                   font_name="Arial", font_size=20,
                   font_color="#94A3B8", align="center")
    
    # ==================== SLIDE 2: Font Showcase ====================
    slide2 = gen.add_slide()
    
    # light background
    gen.add_rectangle(slide2, 0, 0, 16, 9, fill_color="#F8FAFC")
    
    # section header with accent bar
    gen.add_rectangle(slide2, 0, 0, 16, 1.2, fill_color="#2563EB")
    gen.add_textbox(slide2, 0.5, 0.25, 15, 0.8,
                   text="Typography Showcase",
                   font_name="Helvetica Neue", font_size=40,
                   font_color="#FFFFFF", bold=True)
    
    # font samples with cards
    fonts_demo = [
        ("Chalkboard SE", "The quick brown fox jumps over the lazy dog", "#2563EB"),
        ("Bradley Hand", "Handwritten style for creative content", "#7C3AED"),
        ("Georgia", "Elegant serif font for formal documents", "#059669"),
        ("Menlo", "Monospace font for code samples", "#DC2626"),
        ("Helvetica Neue", "Clean sans-serif for modern design", "#F59E0B"),
    ]
    
    y_pos = 1.5
    for i, (font_name, sample_text, accent_color) in enumerate(fonts_demo):
        # card background with shadow
        gen.add_rounded_rectangle(slide2, 0.5, y_pos, 15, 1.2,
                                 fill_color="#FFFFFF", 
                                 border_color=accent_color,
                                 border_width=2,
                                 corner_radius=0.15,
                                 shadow=True)
        
        # font name label
        gen.add_textbox(slide2, 0.8, y_pos + 0.15, 4, 0.4,
                       text=font_name,
                       font_name="Arial", font_size=16,
                       font_color=accent_color, bold=True)
        
        # sample text
        gen.add_textbox(slide2, 0.8, y_pos + 0.55, 14, 0.5,
                       text=sample_text,
                       font_name=font_name, font_size=24,
                       font_color="#1F2937")
        
        y_pos += 1.45
    
    # ==================== SLIDE 3: Shapes & Colors ====================
    slide3 = gen.add_slide()
    
    # gradient-like background
    gen.add_rectangle(slide3, 0, 0, 16, 9, fill_color="#0F172A")
    
    # title
    gen.add_textbox(slide3, 0.5, 0.3, 15, 1,
                   text="Shapes & Colors",
                   font_name="Chalkboard SE", font_size=48,
                   font_color="#FFFFFF", bold=True)
    
    # row 1: rounded rectangles with different colors
    colors_row1 = ["#EF4444", "#F59E0B", "#10B981", "#3B82F6", "#8B5CF6"]
    for i, color in enumerate(colors_row1):
        gen.add_rounded_rectangle(slide3, 0.8 + i * 3, 1.5, 2.5, 1.5,
                                 fill_color=color,
                                 corner_radius=0.25,
                                 shadow=True)
        gen.add_textbox(slide3, 0.8 + i * 3, 1.9, 2.5, 0.6,
                       text=color,
                       font_name="Menlo", font_size=14,
                       font_color="#FFFFFF", align="center",
                       vertical_align="middle")
    
    # row 2: circles
    colors_row2 = ["#EC4899", "#14B8A6", "#F97316", "#6366F1", "#84CC16"]
    for i, color in enumerate(colors_row2):
        gen.add_oval(slide3, 1.1 + i * 3, 3.5, 1.8, 1.8,
                    fill_color=color, shadow=True)
    
    # row 3: mixed shapes with borders
    gen.add_rounded_rectangle(slide3, 1, 5.8, 4, 2.5,
                             fill_color="#1E293B",
                             border_color="#60A5FA",
                             border_width=3,
                             corner_radius=0.3,
                             shadow=True)
    gen.add_textbox(slide3, 1.2, 6.3, 3.6, 1.5,
                   text="Bordered\nRounded Box",
                   font_name="Helvetica Neue", font_size=20,
                   font_color="#60A5FA", align="center")
    
    gen.add_rectangle(slide3, 5.5, 5.8, 4, 2.5,
                     fill_color="#7C3AED",
                     shadow=True)
    gen.add_textbox(slide3, 5.7, 6.3, 3.6, 1.5,
                   text="Rectangle\nwith Shadow",
                   font_name="Helvetica Neue", font_size=20,
                   font_color="#FFFFFF", align="center")
    
    gen.add_oval(slide3, 10.5, 5.8, 2.5, 2.5,
                fill_color="#FBBF24",
                border_color="#FFFFFF",
                border_width=4,
                shadow=True)
    gen.add_textbox(slide3, 10.5, 6.7, 2.5, 0.6,
                   text="Circle",
                   font_name="Arial", font_size=18,
                   font_color="#1F2937", align="center")
    
    gen.add_rounded_rectangle(slide3, 13.5, 5.8, 2, 2.5,
                             fill_color="#F43F5E",
                             corner_radius=0.5,
                             shadow=True)
    
    # ==================== SLIDE 4: Image & Content ====================
    slide4 = gen.add_slide()
    
    # background
    gen.add_rectangle(slide4, 0, 0, 16, 9, fill_color="#FFFFFF")
    
    # header bar
    gen.add_rectangle(slide4, 0, 0, 16, 1.5, fill_color="#1E3A5F")
    gen.add_textbox(slide4, 0.5, 0.35, 10, 1,
                   text="Image Integration",
                   font_name="Chalkboard SE", font_size=44,
                   font_color="#FFFFFF", bold=True)
    
    # left: image placeholder with border
    gen.add_rounded_rectangle(slide4, 0.5, 2, 7, 6,
                             fill_color="#F1F5F9",
                             border_color="#CBD5E1",
                             border_width=2,
                             corner_radius=0.2,
                             shadow=True)
    
    # try to add actual image
    logo_path = Path(__file__).parent.parent / "assets" / "paper2slides_logo.png"
    if logo_path.exists():
        gen.add_image(slide4, str(logo_path), 1.5, 3, width=5)
        gen.add_textbox(slide4, 0.5, 7.2, 7, 0.6,
                       text="paper2slides_logo.png",
                       font_name="Menlo", font_size=14,
                       font_color="#64748B", align="center")
    else:
        gen.add_textbox(slide4, 0.5, 4.5, 7, 1,
                       text="📷 Image Placeholder",
                       font_name="Arial", font_size=28,
                       font_color="#94A3B8", align="center")
    
    # right: content cards
    gen.add_rounded_rectangle(slide4, 8, 2, 7.5, 2,
                             fill_color="#DBEAFE",
                             corner_radius=0.15,
                             shadow=True)
    gen.add_textbox(slide4, 8.3, 2.3, 6.9, 0.5,
                   text="Feature 1: Rich Typography",
                   font_name="Helvetica Neue", font_size=20,
                   font_color="#1E40AF", bold=True)
    gen.add_textbox(slide4, 8.3, 2.9, 6.9, 0.8,
                   text="Support for multiple fonts including Chalkboard SE, Georgia, and custom typefaces.",
                   font_name="Arial", font_size=16,
                   font_color="#3B82F6")
    
    gen.add_rounded_rectangle(slide4, 8, 4.2, 7.5, 2,
                             fill_color="#D1FAE5",
                             corner_radius=0.15,
                             shadow=True)
    gen.add_textbox(slide4, 8.3, 4.5, 6.9, 0.5,
                   text="Feature 2: Shape Styling",
                   font_name="Helvetica Neue", font_size=20,
                   font_color="#047857", bold=True)
    gen.add_textbox(slide4, 8.3, 5.1, 6.9, 0.8,
                   text="Rounded corners, shadows, and custom borders for professional layouts.",
                   font_name="Arial", font_size=16,
                   font_color="#10B981")
    
    gen.add_rounded_rectangle(slide4, 8, 6.4, 7.5, 2,
                             fill_color="#FEE2E2",
                             corner_radius=0.15,
                             shadow=True)
    gen.add_textbox(slide4, 8.3, 6.7, 6.9, 0.5,
                   text="Feature 3: Export Options",
                   font_name="Helvetica Neue", font_size=20,
                   font_color="#B91C1C", bold=True)
    gen.add_textbox(slide4, 8.3, 7.3, 6.9, 0.8,
                   text="Save as PPTX and convert to PDF using LibreOffice.",
                   font_name="Arial", font_size=16,
                   font_color="#EF4444")
    
    # ==================== SLIDE 5: Multi-format Text ====================
    slide5 = gen.add_slide()
    
    # background
    gen.add_rectangle(slide5, 0, 0, 16, 9, fill_color="#FAFAFA")
    
    # header
    gen.add_rounded_rectangle(slide5, 0.3, 0.3, 15.4, 1.4,
                             fill_color="#7C3AED",
                             corner_radius=0.2,
                             shadow=True)
    gen.add_textbox(slide5, 0.5, 0.5, 15, 1,
                   text="Mixed Content Layout",
                   font_name="Chalkboard SE", font_size=44,
                   font_color="#FFFFFF", bold=True)
    
    # left panel with multiline text
    gen.add_rounded_rectangle(slide5, 0.5, 2, 7, 6.5,
                             fill_color="#FFFFFF",
                             border_color="#E5E7EB",
                             border_width=1,
                             corner_radius=0.2,
                             shadow=True)
    
    gen.add_multiline_text(slide5, 0.8, 2.3, 6.4, 5.5, [
        {"text": "Key Points", "font_name": "Helvetica Neue", "font_size": 28, 
         "font_color": "#1F2937", "bold": True},
        {"text": "", "font_size": 12},
        {"text": "• Flexible text formatting", "font_name": "Arial", "font_size": 20, 
         "font_color": "#374151"},
        {"text": "• Multiple font families", "font_name": "Arial", "font_size": 20, 
         "font_color": "#374151"},
        {"text": "• Rich color palette", "font_name": "Arial", "font_size": 20, 
         "font_color": "#374151"},
        {"text": "• Shape customization", "font_name": "Arial", "font_size": 20, 
         "font_color": "#374151"},
        {"text": "", "font_size": 12},
        {"text": "Handwritten note style:", "font_name": "Arial", "font_size": 16, 
         "font_color": "#6B7280", "italic": True},
        {"text": "This looks like a hand note!", "font_name": "Bradley Hand", "font_size": 22, 
         "font_color": "#7C3AED"},
    ])
    
    # right panel with code-like content
    gen.add_rounded_rectangle(slide5, 8, 2, 7.5, 6.5,
                             fill_color="#1E293B",
                             corner_radius=0.2,
                             shadow=True)
    
    gen.add_textbox(slide5, 8.3, 2.2, 6.9, 0.6,
                   text="Code Example",
                   font_name="Helvetica Neue", font_size=20,
                   font_color="#94A3B8", bold=True)
    
    gen.add_multiline_text(slide5, 8.3, 2.9, 6.9, 5, [
        {"text": "from pptx import Presentation", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#F472B6"},
        {"text": "from pptx.util import Inches, Pt", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#F472B6"},
        {"text": "", "font_size": 10},
        {"text": "# Create presentation", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#6B7280", "italic": True},
        {"text": "prs = Presentation()", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#A5F3FC"},
        {"text": "slide = prs.slides.add_slide(...)", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#A5F3FC"},
        {"text": "", "font_size": 10},
        {"text": "# Add textbox", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#6B7280", "italic": True},
        {"text": "tb = slide.shapes.add_textbox(", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#A5F3FC"},
        {"text": "    Inches(1), Inches(1),", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#FDE68A"},
        {"text": "    Inches(4), Inches(2)", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#FDE68A"},
        {"text": ")", "font_name": "Menlo", "font_size": 14, 
         "font_color": "#A5F3FC"},
    ])
    
    # ==================== Save and Convert ====================
    output_dir = Path(__file__).parent
    pptx_path = output_dir / "demo_presentation.pptx"
    
    gen.save(str(pptx_path))
    
    # convert to PDF
    pdf_path = gen.convert_to_pdf(str(pptx_path))
    
    return str(pptx_path), pdf_path


if __name__ == "__main__":
    print("=" * 60)
    print("PPT Generator Demo")
    print("=" * 60)
    
    pptx_path, pdf_path = create_demo_presentation()
    
    print("\n" + "=" * 60)
    print("Generation Complete!")
    print(f"PPTX: {pptx_path}")
    if pdf_path:
        print(f"PDF:  {pdf_path}")
    else:
        print("PDF:  Not generated (LibreOffice not available)")
    print("=" * 60)

